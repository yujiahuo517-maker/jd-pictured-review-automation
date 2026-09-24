# -*- coding:utf-8 -*-
from openpyxl import load_workbook, Workbook

from .base_handler import SheetLikeHandler, parse_coord
from ..cell_style import CellStyle, NoneStyle, CellStyleCopier
from ..header import Header, ZeroHeader
from ..tools import process_content_xlsx, get_real_row, get_first_dict_xlsx, get_key_cols


class XLSXHandler(SheetLikeHandler):
    def __init__(self, recorder):
        super().__init__(recorder)
        self.type = 'xlsx'
        self._header = {None: None}
        self._header_row = {None: 1}
        self._data_col = 1
        self._table = None
        self._None_header_is_newest = None
        self._None_header_row_is_newest = None
        self._methods = {'img': img2ws,
                         'link': link2ws,
                         'style': styles2ws,
                         'height': height2ws,
                         'width': width2ws,
                         'data': data2ws}

    @property
    def table(self):
        return self._table

    @property
    def tables(self):
        if not self.path:
            raise RuntimeError('未指定文件路径。')
        return get_tables(self.path)

    @property
    def header_row(self):
        return self._header_row.get(self._table, 1)

    @property
    def data_col(self):
        return self._data_col

    @property
    def follow_styles(self):
        return self._recorder._settings.get('new_row_styles', False)

    @property
    def new_row_height(self):
        return self._recorder._settings.get('new_row_height', None)

    @property
    def new_row_styles(self):
        return self._recorder._settings.get('new_row_styles', None)

    @property
    def link_style(self):
        return self._recorder._settings.get('link_style', None)

    def add_data(self, data, coord=None, table=None):
        coord = self._parse_coord(coord)
        data = self._handle_data(data)
        if (self.data.get(table, None) and coord == self.data[table][-1]['coord']
                and self.data[table][-1]['type'] == 'data'):
            self.data[table][-1]['data'].extend(data)
        else:
            self.data.setdefault(table, []).append({'type': 'data', 'data': data, 'coord': coord})

    def add_link(self, link, coord, content=None, table=None):
        self.data.setdefault(table, []).append({'type': 'link', 'link': link, 'content': content,
                                                'coord': self._parse_coord(coord)})

    def add_img(self, img_path, coord, width=None, height=None, table=None):
        self.data.setdefault(table, []).append({'type': 'img', 'imgPath': img_path, 'width': width, 'height': height,
                                                'coord': self._parse_coord(coord)})

    def add_styles(self, styles, coord=None, rows=None, cols=None, replace=True, table=None):
        self.data.setdefault(table, []).append({'type': 'style', 'mode': 'replace' if replace else 'cover',
                                                'styles': styles, 'coord': (1, 1), 'real_coord': coord, 'rows': rows,
                                                'cols': cols})

    def add_rows_height(self, height, rows=True, table=None):
        self.data.setdefault(table, []).append({'type': 'height', 'rows': rows, 'height': height})

    def add_cols_width(self, width, cols=True, table=None):
        self.data.setdefault(table, []).append({'type': 'width', 'cols': cols, 'width': width})

    def record(self):
        wb, new_file = get_wb(self._recorder)
        tables = wb.sheetnames
        rewrite_method = 'make_num_dict_rewrite' if self.auto_new_header else 'make_num_dict'

        for table, data in self.data.items():
            ws, new_sheet = get_ws(wb, table, tables, new_file)
            new_file = False
            if table is None:
                if self._None_header_is_newest or ws.title not in self._header:
                    self._header[ws.title] = self._header[None]
                    self._None_header_is_newest = None
                if self._None_header_row_is_newest or ws.title not in self._header_row:
                    self._header_row[ws.title] = self._header_row[None]

            begin_row = True
            if new_sheet:
                begin_row = handle_new_sheet(self, ws, data)
            elif self._header.get(ws.title, None) is None:
                self._header[ws.title] = (Header([c.value for c in ws[self._header_row[ws.title]]])
                                          if ws.title in self._header_row else Header())

            header = self._header[ws.title]
            rewrite = False
            if not begin_row and not data[0]['coord'][0]:  # 首行为空，将数据填入首行
                cur = data[0]
                rewrite = self._methods[cur['type']](
                    **{'handler': self,
                       'ws': ws,
                       'data': cur,
                       'coord': (1, header._get_num(cur.get('coord', (1, 1))[1])),
                       'new_row': not cur.get('coord', (1, 1))[0],
                       'header': header,
                       'rewrite': rewrite,
                       'rewrite_method': rewrite_method})
                data = data[1:]

            for cur in data:
                rewrite = self._methods[cur['type']](
                    **{'handler': self,
                       'ws': ws,
                       'data': cur,
                       'coord': get_ws_real_coord(cur.get('coord', (1, 1)), ws, header),
                       'new_row': not cur.get('coord', (1, 1))[0],
                       'header': header,
                       'rewrite': rewrite,
                       'rewrite_method': rewrite_method})

            if rewrite:
                for c in range(1, ws.max_column + 1):
                    ws.cell(self._header_row[ws.title], c, value=header[c])

        wb.save(self.path)
        wb.close()

    def rows(self, cols=True, sign_col=True, signs=None, deny_sign=False, count=None, begin_row=None, end_row=None):
        wb = load_workbook(self.path, data_only=True, read_only=True)
        if self._table and self._table not in [i.title for i in wb.worksheets]:
            raise RuntimeError(f'xlsx文件未包含指定工作表：{self._table}')
        ws = wb[self._table] if self._table else wb.active
        if ws.max_column is None:  # 遇到过read_only时无法获取列数的文件
            wb.close()
            wb = load_workbook(self.path, data_only=True)
            ws = wb[self._table] if self._table else wb.active
        header = self._get_header(ws)
        if sign_col is not True:
            sign_col = header.get_num(sign_col) or 1
        cols = get_key_cols(cols, header)
        if not begin_row:
            begin_row = self._header_row.get(self._table, self._header_row[None]) + 1
        return get_xlsx_rows(header=header, cols=cols, begin_row=begin_row, end_row=end_row or 0,
                             sign_col=sign_col, sign=signs, deny_sign=deny_sign, count=count, ws=ws)

    def clear(self):
        self.data = {}
        self.data_count = 0

    def set_header(self, header, to_file=True, row=None, table=None):
        if not isinstance(header, (list, tuple)):
            raise ValueError('header必须为list或tuple格式。')

        self._recorder.record()
        row = row or self._header_row.get(table, 1)
        with self._recorder._lock:
            header = Header(header)
            if table is None:
                table = self._recorder.table
            elif table is True:
                table = None
            elif not isinstance(table, str):
                raise ValueError('table只能是None、True或str。')
            self._None_header_is_newest = table is None
            self._header[table] = header
            if to_file:
                set_xlsx_header(self._recorder, header, table, row)

        return self._recorder._setter

    def set_header_row(self, num, table=None):
        if num < 0:
            raise ValueError('num不能小于0。')
        self._recorder.record()
        with self._recorder._lock:
            if table is None:
                table = self._table
            elif table is True:
                table = None
            elif not isinstance(table, str):
                raise ValueError('table只能是None、True或str。')
            self._header_row[table] = num
            self._header[table] = ZeroHeader() if num == 0 else None
            self._None_header_is_newest = table is None
            self._None_header_row_is_newest = table is None
        return self._recorder._setter

    def set_data_col(self, col):
        if col is None:
            self._data_col = 0
        elif not isinstance(col, (int, str)) or col == '':
            raise TypeError('col值只能是int、None或非空str。')
        else:
            self._data_col = col
        return self._recorder._setter

    def set_table(self, name):
        self._table = name if name is not True else None
        return self._recorder._setter

    def set_new_row_height(self, height):
        self._recorder._settings['new_row_height'] = height
        if height is not None:
            self._recorder._settings['follow_styles'] = False
            self._methods['data'] = data2ws_style
        else:
            self._methods['data'] = data2ws
        return self._recorder._setter

    def set_new_row_styles(self, styles):
        self._recorder.record()
        self._recorder._settings['new_row_styles'] = styles
        if styles is not None:
            self._recorder._settings['follow_styles'] = False
            self._methods['data'] = data2ws_style
        else:
            self._methods['data'] = data2ws
        return self._recorder._setter

    def set_follow_styles(self, on_off=True):
        self._recorder._settings['follow_styles'] = on_off
        if on_off:
            self._recorder._settings['styles'] = None
            self._recorder._settings['new_row_height'] = None
            self._methods['data'] = data2ws_follow
        else:
            self._methods['data'] = data2ws
        return self._recorder._setter

    def set_link_style(self, style=True):
        if style is True:
            style = CellStyle()
            style.font.set_color("0000FF")
            style.font.set_underline('single')
        self._recorder._settings['link_style'] = style
        return self._recorder._setter

    def _get_header(self, ws=None):
        header = self._header.get(self._table, None)
        if header is not None:
            return header
        if not self.path or not self._recorder._path.exists():
            return None

        if not ws:
            wb = load_workbook(self.path)
            if not self.table:
                ws = wb.active
            elif self.table not in wb.sheetnames:
                wb.close()
                return Header()
            else:
                ws = wb[self.table]
        else:
            wb = None
        header_row = self._header_row.get(self.table, self._header_row[None])
        if header_row > ws.max_row:
            self._header[self.table] = Header()
        else:
            self._header[self.table] = Header([i.value for i in ws[self._header_row.get(self.table,
                                                                                        self._header_row[None])]])

        if wb:
            wb.close()
        return self._header[self.table]


def set_xlsx_header(recorder, header, table, row):
    if not recorder.path:
        raise FileNotFoundError('未指定文件。')
    if recorder._file_exists or recorder._path.exists():
        wb = load_workbook(recorder.path)
        if table:
            ws = wb[table] if table in [i.title for i in wb.worksheets] else wb.create_sheet(title=table)
        else:
            ws = wb.active

    else:
        recorder._path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        if table:
            ws.title = table

    for c, i in header.items():
        ws.cell(row, c, value=process_content_xlsx(i))
    len_row = len(ws[row])
    len_header = len(header)
    if len_row > len_header:
        for c in range(len_header + 1, len_row + 1):
            ws.cell(row, c, value=None)

    wb.save(recorder.path)
    wb.close()
    recorder._file_exists = True


def line2ws(ws, header, row, col, data, rewrite_method, rewrite):
    if isinstance(data, dict):
        data, rewrite, header_len = header.__getattribute__(rewrite_method)(data, 'xlsx', rewrite)
        for c, val in data.items():
            ws.cell(row, c, value=process_content_xlsx(val))
    else:
        for key, val in enumerate(data):
            ws.cell(row, col + key, value=process_content_xlsx(val))
    return rewrite


def line2ws_follow(ws, header, row, col, data, rewrite_method, rewrite, styles, height, new_row):
    if new_row:
        styles2new_row(ws, styles.values(), height, row)

    if isinstance(data, dict):
        data, rewrite, header_len = header.__getattribute__(rewrite_method)(data, 'xlsx', rewrite)
        if new_row:
            for c, val in data.items():
                ws.cell(row, c, value=process_content_xlsx(val))
        else:
            for c, val in data.items():
                styles.get(c, NoneStyle()).to_cell(ws.cell(row, c, value=process_content_xlsx(val)))
    else:
        if new_row:
            for key, val in enumerate(data):
                ws.cell(row, col + key, value=process_content_xlsx(val))
        else:
            for key, val in enumerate(data):
                col1 = col + key
                styles.get(col1, NoneStyle()).to_cell(ws.cell(row, col1, value=process_content_xlsx(val)))
    return rewrite


def data2ws(handler, ws, data, coord, header, rewrite, rewrite_method, new_row):
    row, col = coord
    for r, d in enumerate(data['data'], row):
        rewrite = line2ws(ws, header, r, col, d, rewrite_method, rewrite)
    return rewrite


def data2ws_follow(handler, ws, data, coord, header, rewrite, rewrite_method, new_row):
    row, col = coord
    if row > 1:
        styles = {ind: CellStyleCopier(cell) for ind, cell in enumerate(ws[row - 1], 1)}
        height = ws.row_dimensions[row - 1].height
        for r, d in enumerate(data['data'], row):
            rewrite = line2ws_follow(ws, header, r, col, d, rewrite_method, rewrite, styles, height, new_row)

    else:
        for r, d in enumerate(data['data'], row):
            rewrite = line2ws(ws, header, r, col, d, rewrite_method, rewrite)

    return rewrite


def data2ws_style(handler, ws, data, coord, header, rewrite, rewrite_method, new_row):
    row, col = coord
    if new_row:
        styles = handler.new_row_styles
        if isinstance(styles, dict):
            styles = header.make_num_dict(styles, None)[0]
            styles = [styles.get(c, None) for c in range(1, ws.max_column + 1)]
        elif isinstance(styles, CellStyle):
            styles = [styles] * ws.max_column
        height = ws.row_dimensions[row].height

        for r, d in enumerate(data['data'], row):
            rewrite = line2ws(ws, header, r, col, d, rewrite_method, rewrite)
            styles2new_row(ws, styles, height, r)

    else:
        for r, d in enumerate(data['data'], row):
            rewrite = line2ws(ws, header, r, col, d, rewrite_method, rewrite)

    return rewrite


def styles2new_row(ws, styles, height, row):
    if height is not None:
        ws.row_dimensions[row].height = height
    if styles:
        for c, s in enumerate(styles, start=1):
            if s:
                s.to_cell(ws.cell(row=row, column=c))


def styles2ws(**kwargs):
    ws = kwargs['ws']
    header = kwargs['header']
    data = kwargs['data']
    styles = data['styles']
    coord = data['real_coord']  # 'A3'、'A1:C3'、(1, 3)、['A1', 'B2', 'C3']
    rows = data['rows']  # 3、'1:3'、[1, 2, 3]
    cols = data['cols']
    mode = data['mode'] == 'replace'
    if not styles:
        styles = [NoneStyle()]
    elif isinstance(styles, CellStyle):
        styles = [styles]

    if isinstance(styles, dict):
        for coord, val in styles.items():
            styles2ws(ws=ws, header=header,
                      data={'styles': val, 'real_coord': coord, 'rows': None, 'cols': None, 'mode': mode})
        return

    if rows:
        if isinstance(rows, int):
            styles_len = len(styles)
            for col, cell in enumerate(ws[header.get_num(rows)]):
                styles[col % styles_len].to_cell(cell, replace=mode)

        elif isinstance(rows, str) and ':' in rows:
            begin, end = rows.split(':', 1)
            try:
                begin = header.get_num(int(begin))
                end = header.get_num(int(end))
            except ValueError:
                raise ValueError('行号必须是数字，现在是：', rows)
            if begin > end:
                begin, end = end, begin
            for i in range(begin, end + 1):
                styles2ws(ws=ws, header=header,
                          data={'styles': styles, 'real_coord': None, 'rows': i, 'cols': None, 'mode': mode})

        elif isinstance(rows, (tuple, list)):
            for i in rows:
                styles2ws(ws=ws, header=header,
                          data={'styles': styles, 'real_coord': None, 'rows': i, 'cols': None, 'mode': mode})

    if cols:
        if isinstance(cols, int):  # 列序号
            styles_len = len(styles)
            for col, cell in enumerate(ws[header.get_col(cols)]):
                styles[col % styles_len].to_cell(cell, replace=mode)

        elif isinstance(cols, str):  # 表头值
            cols = header.get_num(cols)
            if cols:
                styles2ws(ws=ws, header=header,
                          data={'styles': styles, 'real_coord': None, 'rows': None, 'cols': cols, 'mode': mode})

        elif isinstance(cols, tuple) and len(cols) == 2:
            begin, end = cols
            begin = header.get_num(begin)
            end = header.get_num(end)
            if begin > end:
                begin, end = end, begin
            for i in range(begin, end + 1):
                styles2ws(ws=ws, header=header,
                          data={'styles': styles, 'real_coord': None, 'rows': None, 'cols': i, 'mode': mode})

        elif isinstance(cols, (tuple, list)):
            for i in cols:
                styles2ws(ws=ws, header=header,
                          data={'styles': styles, 'real_coord': None, 'rows': None, 'cols': i, 'mode': mode})

    if coord:
        if isinstance(coord, str):
            if ':' in coord:
                begin, end = coord.split(':', 1)
                begin = parse_coord(begin)
                end = parse_coord(end)
                begin = f'{header.get_col(begin[1])}{begin[0]}'
                end = f'{header.get_col(end[1])}{end[0]}'
                styles_len = len(styles)
                for row in ws[f'{begin}:{end}']:
                    for col, cell in enumerate(row):
                        styles[col % styles_len].to_cell(cell, replace=mode)
            else:
                coord = parse_coord(coord)
                coord = f'{header.get_col(coord[1])}{coord[0]}'
                styles[0].to_cell(ws[coord], replace=mode)

        elif isinstance(coord, tuple) and len(coord) == 2:
            coord = parse_coord(coord)
            coord = f'{header.get_col(coord[1])}{coord[0]}'
            styles[0].to_cell(ws[coord], replace=mode)

        elif isinstance(coord, (tuple, list)):
            for i in coord:
                styles2ws(ws=ws, header=header,
                          data={'styles': styles, 'real_coord': i, 'rows': None, 'cols': None, 'mode': mode})


def link2ws(**kwargs):
    handler = kwargs['handler']
    data = kwargs['data']
    cell = kwargs['ws'].cell(*kwargs['coord'])
    has_link = bool(cell.hyperlink)
    cell.hyperlink = data['link']
    if data['content'] is not None:
        cell.value = process_content_xlsx(data['content'])
    if data['link']:
        if handler.link_style:
            handler.link_style.to_cell(cell, replace=False)
    elif has_link:
        NoneStyle().to_cell(cell, replace=False)


def img2ws(**kwargs):
    row, col = kwargs['coord']
    data = kwargs['data']
    ws = kwargs['ws']
    from openpyxl.drawing.image import Image
    img = Image(data['imgPath'])
    width, height = data['width'], data['height']
    if width and height:
        img.width = width
        img.height = height
    elif width:
        img.height = int(img.height * (width / img.width))
        img.width = width
    elif height:
        img.width = int(img.width * (height / img.height))
        img.height = height
    # ws.add_image(img, (row, Header._NUM_KEY[col]))
    ws.add_image(img, f'{Header._NUM_KEY[col]}{row}')


def width2ws(**kwargs):
    # 用int表示列序号，str表示表头值，用tuple设置某列到某列，用list指定每一列，为Ture设置所有列
    cols = kwargs['data']['cols']
    width = kwargs['data']['width']
    ws = kwargs['ws']
    header = kwargs['header']

    if isinstance(width, dict):
        for col, val in width.items():
            width2ws(ws=ws, header=header, data={'cols': col, 'width': val})

    elif isinstance(cols, (int, str)):  # 表头值或列序号
        cols = header.get_col(cols)
        if cols:
            ws.column_dimensions[cols].width = width

    elif isinstance(cols, tuple) and len(cols) == 2:  # 连续多列
        beg, end = cols
        if not beg:
            beg = 1
        if not end:
            end = -1

        beg = header.get_num(beg)
        end = header.get_num(end)

        if beg and end:
            if beg > end:
                beg, end = end, beg
            for c in range(beg, end + 1):
                ws.column_dimensions[header.get_col(c)].width = width

    elif isinstance(cols, list):
        for col in cols:
            width2ws(ws=ws, header=header, data={'cols': col, 'width': width})

    elif cols is True:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[ZeroHeader()[col]].width = width


def height2ws(**kwargs):
    # int表示行号，str为'1:3'格式，dict可指定行独立设置高度，list或tuple指定多行设置同一个高度。True设置所有行
    rows = kwargs['data']['rows']
    height = kwargs['data']['height']
    ws = kwargs['ws']

    if isinstance(height, dict):
        for row, val in height.items():
            height2ws(ws=ws, data={'rows': row, 'height': val})

    elif isinstance(rows, int):
        if rows < 1:
            rows = get_real_row(rows, ws.max_row)
        ws.row_dimensions[rows].height = height

    elif isinstance(rows, str) and ':' in rows:
        beg, end = rows.split(':', 1)
        if beg == '':
            beg = 1
        if end == '':
            end = -1
        beg = int(beg)
        end = int(end)
        if beg < 1 or end < 1:
            max_row = ws.max_row
            beg = get_real_row(beg, max_row)
            end = get_real_row(end, max_row)

        if beg > end:
            beg, end = end, beg
        for c in range(beg, end + 1):
            ws.row_dimensions[c].height = height

    elif isinstance(rows, (list, tuple)):
        for row in rows:
            height2ws(ws=ws, data={'rows': row, 'height': height})

    elif rows is True:
        for i in range(1, ws.max_row + 1):
            ws.row_dimensions[i].height = height


def get_wb(recorder):
    if recorder._file_exists or recorder._path.exists():
        wb = load_workbook(recorder.path)
        new_file = False
    else:
        wb = Workbook()
        new_file = True
    return wb, new_file


def get_ws(wb, table, tables, new_file):
    new_sheet = new_file
    if table is None:
        ws = wb.active
        if ws.max_row == 1 and ws.max_column == 1 and ws.cell(row=1, column=1).value is None:
            new_sheet = True

    elif table in tables:
        ws = wb[table]
        if ws.max_row == 1 and ws.max_column == 1 and not ws.cell(row=1, column=1).value:
            new_sheet = True

    elif new_file is True:
        ws = wb.active
        tables.remove(ws.title)
        ws.title = table
        tables.append(table)
        new_sheet = True

    else:
        ws = wb.create_sheet(title=table)
        tables.append(table)
        new_sheet = True

    return ws, new_sheet


def handle_new_sheet(handler, ws, data):
    if not handler._header_row:
        return 0

    if handler._header.get(ws.title, None) is not None:
        for c, h in handler._header[ws.title].items():
            ws.cell(row=handler._header_row[ws.title], column=c, value=h)
        begin_row = handler._header_row

    else:
        data = get_first_dict_xlsx(data)
        if data:
            header = Header([h for h in data.keys() if isinstance(h, str)])
            handler._header[ws.title] = header
            for c, h in header.items():
                ws.cell(row=handler._header_row[ws.title], column=c, value=h)
            begin_row = handler._header_row

        else:
            handler._header[ws.title] = Header()
            begin_row = 0

    return begin_row


def get_ws_real_coord(coord, ws, header):
    row, col = coord
    if row <= 0:
        row = ws.max_row + row + 1
    return 1 if row < 1 else row, header._get_num(col)


def get_xlsx_rows(header, cols, begin_row, end_row, sign_col, sign, deny_sign, count, ws):
    rows = ws.rows
    try:
        for _ in range(begin_row - 1):
            next(rows)
    except StopIteration:
        return []

    if sign_col is True or sign_col > ws.max_column:  # 获取所有行
        if count or end_row:
            rows = list(rows)[:(min(count, end_row - begin_row + 1)
                                if count and end_row else (count or end_row - begin_row + 1))]

        if cols is True:  # 获取整行
            res = [header.make_row_data(ind, {col: cell.value for col, cell in enumerate(row, 1)})
                   for ind, row in enumerate(rows, begin_row)]
        else:  # 只获取对应的列
            res = [header.make_row_data(ind, {col: row[col - 1].value for col in cols})
                   for ind, row in enumerate(rows, begin_row)]

    else:  # 获取符合条件的行
        if count:
            res = get_xlsx_rows_with_count(cols, deny_sign, header, rows,
                                           begin_row, end_row, sign_col, sign, count)
        else:
            res = get_xlsx_rows_without_count(cols, deny_sign, header, rows, begin_row, end_row,
                                              sign_col, sign)

    ws.parent.close()
    return res


def get_xlsx_rows_with_count(cols, deny_sign, header, rows, begin_row, end_row, sign_col, sign, count):
    got = 0
    res = []
    if cols is True:  # 获取整行
        if deny_sign:
            for ind, row in enumerate(rows, begin_row):
                if got == count or (end_row and ind > end_row):
                    break
                if row[sign_col - 1].value not in sign:
                    res.append(header.make_row_data(ind, {col: cell.value for col, cell in enumerate(row, 1)}))
                    got += 1
        else:
            for ind, row in enumerate(rows, begin_row):
                if got == count or (end_row and ind > end_row):
                    break
                if row[sign_col - 1].value in sign:
                    res.append(header.make_row_data(ind, {col: cell.value for col, cell in enumerate(row, 1)}))
                    got += 1

    else:  # 只获取对应的列
        if deny_sign:
            for ind, row in enumerate(rows, begin_row):
                if got == count or (end_row and ind > end_row):
                    break
                if row[sign_col - 1].value not in sign:
                    res.append(header.make_row_data(ind, {col: row[col - 1].value for col in cols}))
                    got += 1
        else:
            for ind, row in enumerate(rows, begin_row):
                if got == count or (end_row and ind > end_row):
                    break
                if row[sign_col - 1].value in sign:
                    res.append(header.make_row_data(ind, {col: row[col - 1].value for col in cols}))
                    got += 1
    return res


def get_xlsx_rows_without_count(cols, deny_sign, header, rows, begin_row, end_row, sign_col, sign):
    if end_row:
        if end_row < begin_row:
            return []
        rows = list(rows)[:end_row - begin_row + 1]
    if cols is True:  # 获取整行
        if deny_sign:
            return [header.make_row_data(ind, {col: cell.value for col, cell in enumerate(row, 1)})
                    for ind, row in enumerate(rows, begin_row)
                    if row[sign_col - 1].value not in sign]
        else:
            return [header.make_row_data(ind, {col: cell.value for col, cell in enumerate(row, 1)})
                    for ind, row in enumerate(rows, begin_row)
                    if row[sign_col - 1].value in sign]

    else:  # 只获取对应的列
        if deny_sign:
            return [header.make_row_data(ind, {col: row[col - 1].value for col in cols})
                    for ind, row in enumerate(rows, begin_row)
                    if row[sign_col - 1].value not in sign]
        else:
            return [header.make_row_data(ind, {col: row[col - 1].value for col in cols})
                    for ind, row in enumerate(rows, begin_row)
                    if row[sign_col - 1].value in sign]


def get_tables(path):
    wb = load_workbook(path)
    tables = wb.sheetnames
    wb.close()
    return tables
