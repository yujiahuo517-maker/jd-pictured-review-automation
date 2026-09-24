from __future__ import annotations
import argparse, hashlib, json, os, re, sys, time
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook

SKILL = Path(__file__).resolve().parents[1]
WORK = Path(os.environ.get("JD_PIC_WORK_DIR") or (Path.home() / "Desktop" / "jd-pictured-review-run"))
XLSX = Path(os.environ.get("JD_PIC_SPU_XLSX") or (Path.home() / "Desktop" / "SPU.xlsx"))
TODAY = time.strftime("%Y%m%d")
DISPLAY_LIMIT = int(os.environ.get("JD_PIC_DISPLAY_LIMIT", "5"))
SPU_TOTAL_LIMIT = int(os.environ.get("JD_PIC_SPU_LIMIT", "100"))
TEXT_ONLY_LIMIT = max(0, SPU_TOTAL_LIMIT - DISPLAY_LIMIT)
WORK.mkdir(parents=True, exist_ok=True)

sys.path.insert(0, str(SKILL / "vendor" / "blacklight" / "src"))
sys.path.insert(0, str(SKILL / "runtime" / "python-libs"))
os.environ.setdefault("BLACKLIGHT_HOME", str(SKILL / "runtime" / "state"))

from blacklight.llm import gateway
from blacklight.pic import client, imgzone, screen

STYLE_PROFILES = [
    "挂拍：衣架挂在浅色墙面或衣柜旁，自然室内光，完整展示衣服正面，手机随手拍质感。",
    "平铺：商品平铺在床面或桌面，旁边可有日常物件，俯拍构图，背景干净但不商业棚拍。",
    "上身镜拍：真人穿着对镜自拍，脸部可被手机遮挡，居家卧室/衣柜背景，保持真实买家晒单感。",
    "局部细节：近距离拍衣料纹理、领口/袖口/拉链等细节，背景为桌面或床面，画面自然。",
    "使用场景：居家、走廊、阳台或轻运动后随手拍，构图与前几张明显不同。",
]

TEXT_ANGLES = [
    "先说上身感，再说面料，最后轻带尺码或颜色。",
    "先说面料手感，再说日常搭配或使用场景，避免固定开头。",
    "先提颜色/尺码匹配，再描述版型和舒适度。",
    "用一句生活化口吻开头，再补充做工或细节感受。",
    "先说整体满意点，再说一个具体细节，语气自然不夸张。",
    "从试穿/使用场景切入，句式与前后SKU错开。",
]

FORBIDDEN_TEXT_TERMS = (
    "物流", "快递", "客服", "售后", "正品", "品牌授权", "疗效", "最", "第一",
)


def read_spus() -> list[str]:
    wb = load_workbook(XLSX, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                text = str(value or "").strip()
                if text.isdigit() and len(text) >= 6 and text not in out:
                    out.append(text)
    return out


def image_key(url: str) -> str:
    path = urlparse(url or "").path.lower()
    return path.rsplit("/", 1)[-1] if path else str(url).lower()


def sku_parts(name: str) -> tuple[str, str]:
    color = next((c for c in ["纯黑", "纯蓝", "纯灰", "纯白", "黑色", "蓝色", "灰色", "白色", "军绿色", "卡其色", "杏色", "藏青", "米色"] if c in name), "")
    size = ""
    match = re.search(r"(?<![A-Za-z0-9])(L|XL|2XL|3XL|4XL|5XL|6XL|M|S|170|175|180|185|190)(?![A-Za-z0-9])", name, re.I)
    if match:
        size = match.group(1).upper()
    return color, size


def stable_index(*parts: str, modulo: int) -> int:
    seed = "|".join(parts).encode("utf-8", errors="ignore")
    return int(hashlib.md5(seed).hexdigest()[:8], 16) % modulo


def fallback_text(name: str, variant: int = 0) -> str:
    color, size = sku_parts(name)
    color_part = f"{color}" if color else "这款"
    size_part = f"，{size}尺码上身比较合适" if size else "，尺码选择起来比较省心"
    size_start = f"{size}码试穿下来比较合适" if size else "尺码感觉比较合适"
    if "背心" in name or "无袖" in name:
        options = [
            f"{color_part}背心穿着比较舒服，面料轻薄透气，夏天运动和日常居家穿都合适{size_part}。",
            f"面料摸着挺清爽，{color_part}颜色日常好搭，居家或运动时穿着都不闷{size_part}。",
            f"{size_start}，肩口位置不勒，整体比较轻便，热天单穿也挺自在。",
            f"拿到后试了下，版型不会太贴身，{color_part}背心搭短裤长裤都方便。",
        ]
        return options[variant % len(options)]
    if "外套" in name or "冲锋衣" in name or "夹克" in name:
        options = [
            f"{color_part}外套上身版型不错，日常通勤和户外活动都合适，面料手感舒适{size_part}。",
            f"面料摸着比想象中扎实，穿起来不显笨重，{color_part}颜色也比较耐看{size_part}。",
            f"{size_start}，袖口和走线看着规整，平时出门搭牛仔裤挺省心。",
            f"日常外出穿着挺方便，版型利落不拖沓，活动时没有明显束缚感。",
            f"衣服厚度够用，领口和拉链细节处理还可以，整体是实穿型。",
        ]
        return options[variant % len(options)]
    if "裤" in name:
        options = [
            f"{color_part}裤子穿着比较舒服，版型日常好搭，活动起来不紧绷{size_part}。",
            f"裤型比较利落，面料手感舒服，坐下和走动都没有明显拘束感{size_part}。",
            f"{size_start}，颜色和上衣很好配，日常通勤或者休闲穿都合适。",
            f"试穿后感觉腰腿位置处理得不错，整体不臃肿，做工也比较稳。",
        ]
        return options[variant % len(options)]
    options = [
        f"这款商品实物和描述基本一致，做工看着不错，日常使用比较方便{size_part}。",
        f"整体质感比预期稳，细节处理比较规整，平时用起来挺顺手{size_part}。",
        f"到手先看了做工，边角处理还可以，实际使用感受比较自然。",
        f"款式属于耐看型，搭配和使用都不费劲，整体体验符合预期。",
    ]
    return options[variant % len(options)]


def sanitize_text(text: str, fallback: str) -> str:
    cleaned = re.sub(r"\s+", "", (text or "").strip())
    if not cleaned or len(cleaned) < 8 or len(cleaned) > 120:
        return fallback
    if any(term in cleaned for term in FORBIDDEN_TEXT_TERMS):
        return fallback
    return cleaned


def slot_key(item: dict) -> str:
    return f"{item['sku_id']}#{int(item.get('slot_index') or 0)}"


def generate_texts(items: list[dict], spu: str) -> dict[str, str]:
    texts = {
        slot_key(it): fallback_text(
            it.get("sku_name", ""),
            stable_index(spu, it["sku_id"], str(it.get("slot_index", 0)), modulo=97),
        )
        for it in items
    }
    for chunk_index, start in enumerate(range(0, len(items), 25), 1):
        chunk = items[start:start + 25]
        payload = [
            {
                "row_id": slot_key(it),
                "sku_id": it["sku_id"],
                "sku_name": it.get("sku_name", ""),
                "review_no": int(it.get("slot_index") or 0) + 1,
                "writing_angle": TEXT_ANGLES[(start + idx) % len(TEXT_ANGLES)],
            }
            for idx, it in enumerate(chunk)
        ]
        by_key = {slot_key(it): it for it in chunk}
        prompt = (
            "请根据SKU名称生成自然买家评价文案。要求：中文25到60字；不编造物流、售后、品牌授权、疗效、极限效果；"
            "不出现差评、真假质疑；可围绕颜色、尺码、穿着舒适、面料手感、日常使用体验表达。"
            "同一SKU可能需要多条评价，同一批文案在符合产品事实的前提下必须明显错开开头、句式、语序和体验角度，"
            "避免连续使用‘这款/这件/面料/尺码’等相同开头，不要机械重复‘穿着舒服+面料不错+尺码合适’结构；"
            "每条按writing_angle组织表达，并用row_id逐条返回。返回严格JSON数组，每项包含row_id和text。SKU列表："
            + json.dumps(payload, ensure_ascii=False)
        )
        try:
            arr = gateway.chat_json(prompt, timeout=120)
            if isinstance(arr, dict):
                arr = arr.get("items") or arr.get("rows") or []
            for row in arr or []:
                key = str(row.get("row_id") or "").strip()
                text = str(row.get("text") or row.get("eval_content") or "").strip()
                if key in texts:
                    item = by_key.get(key) or {}
                    fallback = fallback_text(
                        item.get("sku_name", ""),
                        stable_index(spu, key, "retry", modulo=101),
                    )
                    texts[key] = sanitize_text(text, fallback)
            print(f"spu={spu} text_chunk={chunk_index}/{(len(items)+24)//25} ok", flush=True)
        except Exception as exc:
            print(f"spu={spu} text_chunk={chunk_index} fallback {type(exc).__name__}: {str(exc)[:120]}", flush=True)

    seen = set()
    for pos, item in enumerate(items):
        key = slot_key(item)
        text = texts[key]
        if text in seen:
            for bump in range(1, 40):
                alt = fallback_text(
                    item.get("sku_name", ""),
                    stable_index(spu, key, str(pos), str(bump), modulo=997) + bump,
                )
                if alt not in seen:
                    text = alt
                    break
        texts[key] = text
        seen.add(text)
    return texts


def choose_display(items: list[dict], limit: int = DISPLAY_LIMIT) -> list[dict]:
    selected, selected_keys, seen_img = [], set(), set()
    for it in items:
        key = image_key(it.get("sku_image", ""))
        skey = slot_key(it)
        if key not in seen_img and skey not in selected_keys:
            selected.append(it); selected_keys.add(skey); seen_img.add(key)
        if len(selected) >= limit:
            return selected
    for it in items:
        skey = slot_key(it)
        if skey not in selected_keys:
            selected.append(it); selected_keys.add(skey)
        if len(selected) >= limit:
            break
    return selected


def text_from_ai(res: dict) -> str:
    return (res.get("eval_content") or (res.get("raw") or {}).get("aiEvalContent") or "").strip()


def platform_ai_text(sku: str, images: list[str]) -> str:
    try:
        return text_from_ai(client.ai_generate(sku, images))
    except Exception:
        return ""


def get_pending_items(spu: str) -> tuple[dict, list[dict]]:
    raw = client.targets(limit=0, page_size=100, spu_ids=[spu], image_filter=0, max_pages=200)
    max_per = int(raw.get("max_per_sku") or 10)
    seen, items = set(), []
    for it in raw.get("items", []):
        sku = it.get("sku_id")
        used = int(it.get("used") or 0)
        if sku and sku not in seen and it.get("sku_image") and used < max_per:
            row = dict(it)
            row["used"] = used
            row["remaining_slots"] = max_per - used
            seen.add(sku); items.append(row)
    return raw, items


def expand_slots(items: list[dict], max_per: int, limit: int = SPU_TOTAL_LIMIT) -> list[dict]:
    slots = []
    max_remaining = max((int(it.get("remaining_slots") or 0) for it in items), default=0)
    for offset in range(max_remaining):
        for it in items:
            used = int(it.get("used") or 0)
            slot_index = used + offset
            if slot_index >= max_per:
                continue
            slot = dict(it)
            slot["slot_index"] = slot_index
            slot["slot_no"] = offset + 1
            slot["slot_key"] = slot_key(slot)
            slots.append(slot)
            if len(slots) >= limit:
                return slots
    return slots


def available_slot_count(items: list[dict], max_per: int) -> int:
    return sum(max(0, max_per - int(it.get("used") or 0)) for it in items)


def submit_text_only(spu_dir: Path, spu: str, slots: list[dict], texts: dict[str, str]) -> list[dict]:
    all_results = []
    print(f"spu={spu} text_slots={len(slots)}", flush=True)
    for n, it in enumerate(slots, 1):
        sku = it["sku_id"]
        key = slot_key(it)
        fallback = fallback_text(it.get("sku_name", ""), stable_index(spu, key, "late", modulo=997))
        body = {
            "skuId": sku,
            "evalContent": sanitize_text(texts.get(key), fallback),
            "images": "",
            "isCheck": False,
            "skuTaskIndex": int(it.get("slot_index") or 0),
        }
        rec = {"sku_id": sku, "slot_index": body["skuTaskIndex"], "eval_content": body["evalContent"]}
        try:
            rec["response"] = client._call(client.FID_IMPORT, body)
            rec["ok"] = True
        except Exception as exc:
            rec["ok"] = False
            rec["error"] = type(exc).__name__ + ": " + str(exc)
        all_results.append(rec)
        if n % 20 == 0 or n == len(slots):
            print(f"spu={spu} text_submitted={n}/{len(slots)}", flush=True)
        time.sleep(0.2)
    (spu_dir / "text_submit_results.json").write_text(json.dumps(all_results, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return all_results


def handle_spu(spu: str) -> dict:
    spu_dir = WORK / spu
    gen_dir = spu_dir / "generated"
    spu_dir.mkdir(parents=True, exist_ok=True)
    gen_dir.mkdir(parents=True, exist_ok=True)

    raw_targets, items = get_pending_items(spu)
    max_per = int(raw_targets.get("max_per_sku") or 10)
    total_available_slots = available_slot_count(items, max_per)
    slots = expand_slots(items, max_per, SPU_TOTAL_LIMIT)
    (spu_dir / "targets_initial.json").write_text(json.dumps(raw_targets, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (spu_dir / "slot_plan.json").write_text(json.dumps(slots, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(
        f"spu={spu} pending_unique={len(items)} available_slots={total_available_slots} "
        f"planned_slots={len(slots)} max_per_sku={max_per} total={raw_targets.get('total')}",
        flush=True,
    )
    if not slots:
        return {"spu": spu, "pending": 0, "message": "no available slots"}

    texts = generate_texts(slots, spu)
    (spu_dir / "text_candidates.json").write_text(json.dumps([
        {
            "row_id": slot_key(it),
            "sku_id": it["sku_id"],
            "slot_index": it.get("slot_index"),
            "sku_name": it.get("sku_name", ""),
            "candidate_text": texts[slot_key(it)],
        }
        for it in slots
    ], ensure_ascii=False, indent=2), encoding="utf-8")

    selected = choose_display(slots, min(DISPLAY_LIMIT, len(slots)))
    selected_keys = {slot_key(x) for x in selected}
    queue = selected + [it for it in slots if slot_key(it) not in selected_keys]
    display_rows, display_steps, attempted = [], [], set()
    for it in queue:
        if len(display_rows) >= DISPLAY_LIMIT:
            break
        sku = it["sku_id"]
        skey = slot_key(it)
        if skey in attempted:
            continue
        attempted.add(skey)
        style = STYLE_PROFILES[len(display_rows) % len(STYLE_PROFILES)]
        stem = f"{sku}_slot{int(it.get('slot_index') or 0):02d}_piceval_{TODAY}_{len(display_steps)+1:02d}"
        prompt = (
            "基于输入商品图，生成一张真实买家晒单风格照片。保持商品外观、颜色、版型一致；"
            "无文字、无水印、无logo篡改、无夸大效果；商品清晰可见。"
            f"本张差异化拍摄要求：{style}"
        )
        step = {"sku_id": sku, "slot_index": int(it.get("slot_index") or 0), "sku_name": it.get("sku_name", ""), "style": style}
        try:
            print(f"spu={spu} image_attempt={len(display_steps)+1} sku={sku} slot={step['slot_index']}", flush=True)
            gen = gateway.image_edit(prompt=prompt, images=[it["sku_image"]], out_dir=str(gen_dir), stem=stem, size="1024x1024", n=1)
            local_path = gen["paths"][0]
            uploaded = imgzone.upload(local_path, file_name=f"{stem}.png", cate_id="0")
            image_url = uploaded["url"]
            alive = imgzone.check_alive(image_url)
            fallback = fallback_text(it.get("sku_name", ""), stable_index(spu, skey, "display", modulo=997))
            text = sanitize_text(texts.get(skey) or platform_ai_text(sku, [image_url]), fallback)
            row = {"sku_id": sku, "sku_name": it.get("sku_name", ""), "eval_content": text, "images": [image_url]}
            screened = screen.screen_rows(rows=[row])
            step.update({"local_path": local_path, "image_url": image_url, "alive": alive, "eval_content": text, "screen": screened})
            if screened.get("ok"):
                display_rows.append(screened["ok"][0])
                step["status"] = "ok"
            else:
                step["status"] = "blocked"
        except Exception as exc:
            step.update({"status": "error", "error": type(exc).__name__ + ": " + str(exc)})
        display_steps.append(step)
        (spu_dir / "display_steps.json").write_text(json.dumps(display_steps, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print(f"spu={spu} display_ok={len(display_rows)} status={step['status']}", flush=True)

    display_submit = {"executed": False, "success_count": 0, "results": [], "rejected": []}
    if display_rows:
        verify = client.targets(limit=0, page_size=100, sku_ids=list({r["sku_id"] for r in display_rows}), image_filter=0, max_pages=20)
        remain = {x["sku_id"] for x in verify.get("items", []) if x.get("result_type") == 0 and int(x.get("used") or 0) < max_per}
        display_rows = [r for r in display_rows if r["sku_id"] in remain]
        if display_rows:
            dry = client.import_rows_dryrun(display_rows)
            (spu_dir / "display_dryrun.json").write_text(json.dumps(dry, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            if dry.get("ok"):
                display_submit = client.import_rows(dry["ok"], confirm=dry["confirm_token"])
    (spu_dir / "display_submit.json").write_text(json.dumps(display_submit, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    display_successes = [r for r in display_submit.get("results", []) if r.get("ok")]
    display_used_slots = {(str(r.get("sku_id")), int(r.get("sku_task_index") or 0)) for r in display_successes}
    print(f"spu={spu} display_submitted={len(display_successes)}", flush=True)

    text_slots = [
        it for it in slots
        if (str(it["sku_id"]), int(it.get("slot_index") or 0)) not in display_used_slots
    ]
    text_capacity = max(0, SPU_TOTAL_LIMIT - DISPLAY_LIMIT)
    text_slots = text_slots[:text_capacity]
    text_results = submit_text_only(spu_dir, spu, text_slots, texts)
    raw_final, final_items = get_pending_items(spu)
    final_max_per = int(raw_final.get("max_per_sku") or max_per)
    final_remaining_slots = available_slot_count(final_items, final_max_per)
    summary = {
        "spu": spu,
        "initial_pending_unique": len(items),
        "initial_available_slots": total_available_slots,
        "planned_total_limit": SPU_TOTAL_LIMIT,
        "planned_total_rows": min(len(slots), DISPLAY_LIMIT + max(0, SPU_TOTAL_LIMIT - DISPLAY_LIMIT)),
        "planned_text_limit": max(0, SPU_TOTAL_LIMIT - DISPLAY_LIMIT),
        "text_candidates": len(slots),
        "display_attempted": len(display_steps),
        "display_submitted": len(display_successes),
        "text_attempted": len(text_results),
        "text_success": sum(1 for r in text_results if r.get("ok")),
        "text_failed": sum(1 for r in text_results if not r.get("ok")),
        "final_pending_sku_count": len(final_items),
        "final_remaining_slots": final_remaining_slots,
        "spu_limit_reached": total_available_slots > SPU_TOTAL_LIMIT,
        "display_sku_slots": sorted(f"{sku}#{idx}" for sku, idx in display_used_slots),
    }
    (spu_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print("SPU_SUMMARY " + json.dumps(summary, ensure_ascii=False), flush=True)
    return summary



def main():
    global XLSX, WORK, DISPLAY_LIMIT, SPU_TOTAL_LIMIT, TEXT_ONLY_LIMIT

    parser = argparse.ArgumentParser(description="按SPU表批量提交京喜带图/纯文案评价")
    parser.add_argument("--xlsx", default=str(XLSX), help="SPU表路径，默认桌面SPU.xlsx")
    parser.add_argument("--work-dir", default=str(WORK), help="产物目录，默认桌面jd-pictured-review-run")
    parser.add_argument("--display-limit", type=int, default=DISPLAY_LIMIT, help="每个SPU带图评价数量，默认5")
    parser.add_argument("--spu-limit", type=int, default=SPU_TOTAL_LIMIT, help="每个SPU总提交条数上限，默认100")
    args = parser.parse_args()

    XLSX = Path(args.xlsx)
    WORK = Path(args.work_dir)
    DISPLAY_LIMIT = int(args.display_limit)
    SPU_TOTAL_LIMIT = int(args.spu_limit)
    TEXT_ONLY_LIMIT = max(0, SPU_TOTAL_LIMIT - DISPLAY_LIMIT)
    WORK.mkdir(parents=True, exist_ok=True)

    if not XLSX.exists():
        raise FileNotFoundError(f"SPU表不存在：{XLSX}")
    spus = read_spus()
    print("xlsx=" + str(XLSX), flush=True)
    print("work_dir=" + str(WORK), flush=True)
    print("display_limit=" + str(DISPLAY_LIMIT), flush=True)
    print("spu_total_limit=" + str(SPU_TOTAL_LIMIT), flush=True)
    print("spus=" + ",".join(spus), flush=True)
    summaries = []
    for spu in spus:
        try:
            summaries.append(handle_spu(spu))
        except Exception as exc:
            summary = {"spu": spu, "error": type(exc).__name__ + ": " + str(exc)}
            summaries.append(summary)
            print("SPU_ERROR " + json.dumps(summary, ensure_ascii=False), flush=True)
    (WORK / "summary.json").write_text(json.dumps(summaries, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print("ALL_SUMMARY " + json.dumps(summaries, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
